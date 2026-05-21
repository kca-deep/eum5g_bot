"""
Convert eum5G xlsx sources into Dify-ready Markdown.

Outputs:
  docs/result/문의로그_2023-2025.md   (488 Q&A chunks)
  docs/result/특화망_구축사례.md       (46 case study chunks)

Each top-level "##" section becomes one chunk in Dify's header-based chunking.
Run:  python scripts/xlsx_to_dify_md.py
"""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
OUT = DOCS / "result"
OUT.mkdir(parents=True, exist_ok=True)

INQUIRY_XLSX = DOCS / "문의로그 DB화 편집본(2023-2025).xlsx"
CASE_XLSX = DOCS / "특화망 구축사례 DB화 파일_기업정보 삭제.xlsx"

CATEGORY_FIX = {
    "네트워크 구축/운영": "네트워크 구축/운용",
    "네트워크 구축/운용 ": "네트워크 구축/운용",
    "5G장비": "5G 장비",
    "5G단말": "5G 단말",
    "기타 ": "기타",
    "현대자동차": "기타",
}
TYPE_FIX = {"구축": "공급"}


def norm_text(v: object) -> str:
    if v is None:
        return ""
    s = str(v).replace("\r\n", "\n").replace("\r", "\n").strip()
    s = re.sub(r"\n{3,}", "\n\n", s)
    return s


def norm_category(v: object) -> str:
    s = norm_text(v)
    return CATEGORY_FIX.get(s, s)


def norm_type(v: object) -> str:
    s = norm_text(v)
    return TYPE_FIX.get(s, s)


def collapse_blank_lines(s: str) -> str:
    """Avoid \n\n inside a chunk body — keeps Dify's paragraph splitter from
    breaking a single Q/A mid-content."""
    return re.sub(r"\n{2,}", "\n", s)


def build_inquiry_md() -> tuple[str, int]:
    wb = openpyxl.load_workbook(INQUIRY_XLSX, data_only=True)
    ws = wb["문의로그"]

    lines: list[str] = [
        "# 5G 특화망 지원센터 문의응답 DB (2023-2025)",
        "",
        "출처: 한국방송통신전파진흥원(KCA) 5G 특화망 지원센터 (local5g@kca.kr)",
        "각 항목은 실제 문의/답변 사례이며, 분야·유형으로 분류되어 있습니다.",
        "",
        "---",
        "",
    ]

    kept = 0
    seq = 0
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row, values_only=True):
        type_raw, cat_raw, q_raw, a_raw = row[0], row[1], row[2], row[3]
        q = norm_text(q_raw)
        a = norm_text(a_raw)
        if not q and not a:
            continue
        seq += 1
        t = norm_type(type_raw) or "미상"
        c = norm_category(cat_raw) or "미상"
        q_body = collapse_blank_lines(q) if q else "(질문 내용 없음)"
        a_body = collapse_blank_lines(a) if a else "(답변 내용 없음)"

        lines.append(f"## 문의 #{seq:03d} | 분야: {c} | 유형: {t}")
        lines.append("")
        lines.append("**질문**")
        lines.append(q_body)
        lines.append("")
        lines.append("**답변**")
        lines.append(a_body)
        lines.append("")
        lines.append("---")
        lines.append("")
        kept += 1

    return "\n".join(lines).rstrip() + "\n", kept


def build_case_md() -> tuple[str, int]:
    wb = openpyxl.load_workbook(CASE_XLSX, data_only=True)
    ws = wb["Sheet1"]

    # Header layout (row 2 + row 3 sub-headers for hardware composition):
    #  A:할당기간  B:TYPE  C:주파수대역  D:망 구축 형태
    #  E:RU  F:DU  G:CU  H:코어  I:단말   (E2:I2 merged = "하드웨어 구성")
    #  J:도입시설 규모  K:초기투자비용  L:산업분류  M:적용서비스
    #  N:핵심 요구사항  O:기존 네트워크 환경  P:도입효과
    #  Q:구축 파트너사  R:구축 소요기간  S:운영 방식
    #  T:도입시 가장 큰 진입장벽  U:기타
    field_order = [
        ("할당기간", 0),
        ("TYPE", 1),
        ("주파수대역", 2),
        ("망 구축 형태", 3),
        # 하드웨어 구성은 별도 처리 (col 4~8)
        ("도입시설 규모", 9),
        ("초기투자비용", 10),
        ("산업분류", 11),
        ("적용서비스", 12),
        ("핵심 요구사항", 13),
        ("기존 네트워크 환경", 14),
        ("도입효과", 15),
        ("구축 파트너사", 16),
        ("구축 소요기간", 17),
        ("운영 방식", 18),
        ("도입시 가장 큰 진입장벽", 19),
        ("기타", 20),
    ]
    hw_labels = [("RU", 4), ("DU", 5), ("CU", 6), ("코어", 7), ("단말", 8)]

    lines: list[str] = [
        "# 이음5G 특화망 구축사례 DB",
        "",
        "출처: 5G 특화망 지원센터. 기업정보(고객사명)는 사전에 삭제된 상태입니다.",
        "각 사례는 할당기간·산업분류·하드웨어 구성·도입효과 등 21개 항목으로 정리되어 있습니다.",
        "",
        "---",
        "",
    ]

    kept = 0
    seq = 0
    # Data starts at row 4
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, values_only=True):
        # Skip rows where all interesting cells are empty
        if not any(norm_text(c) for c in row):
            continue
        seq += 1

        type_v = norm_text(row[1])
        industry_v = norm_text(row[11])
        net_form_v = norm_text(row[3])

        header_bits = [f"사례 #{seq:02d}"]
        if industry_v:
            header_bits.append(f"산업: {industry_v}")
        if net_form_v:
            header_bits.append(f"망 형태: {net_form_v}")
        if type_v:
            header_bits.append(f"TYPE: {type_v}")
        lines.append("## " + " | ".join(header_bits))
        lines.append("")

        for label, idx in field_order:
            val = norm_text(row[idx])
            if not val:
                continue
            if "\n" in val:
                # Multi-line value: render as label + indented block
                indented = "\n".join("  " + ln for ln in val.split("\n"))
                lines.append(f"- **{label}**:")
                lines.append(indented)
            else:
                lines.append(f"- **{label}**: {val}")

        # 하드웨어 구성 sub-block
        hw_vals = [(lbl, norm_text(row[idx])) for lbl, idx in hw_labels]
        hw_vals = [(l, v) for l, v in hw_vals if v]
        if hw_vals:
            lines.append("- **하드웨어 구성**:")
            for lbl, val in hw_vals:
                if "\n" in val:
                    first, *rest = val.split("\n")
                    lines.append(f"  - {lbl}: {first}")
                    for r in rest:
                        lines.append(f"    {r}")
                else:
                    lines.append(f"  - {lbl}: {val}")

        lines.append("")
        lines.append("---")
        lines.append("")
        kept += 1

    return "\n".join(lines).rstrip() + "\n", kept


def main() -> None:
    inquiry_md, n_inq = build_inquiry_md()
    case_md, n_case = build_case_md()

    inq_path = OUT / "문의로그_2023-2025.md"
    case_path = OUT / "특화망_구축사례.md"
    inq_path.write_text(inquiry_md, encoding="utf-8")
    case_path.write_text(case_md, encoding="utf-8")

    print(f"[OK] {inq_path.relative_to(ROOT)}  ({n_inq} Q&A, {len(inquiry_md):,} chars)")
    print(f"[OK] {case_path.relative_to(ROOT)}  ({n_case} cases, {len(case_md):,} chars)")


if __name__ == "__main__":
    main()
